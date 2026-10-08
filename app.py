
import os, re, sqlite3, secrets, calendar
from datetime import datetime
from functools import wraps
from pathlib import Path
from flask import Flask, render_template, request, jsonify, session, send_from_directory
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import load_workbook

BASE=Path(__file__).resolve().parent
DB=BASE/"schoolpulse.db"
UPLOADS=Path(os.environ.get("EDUTRACK_UPLOADS", str(BASE/"uploads")))
UPLOADS.mkdir(parents=True, exist_ok=True)

app=Flask(__name__)
app.secret_key=os.environ.get("SCHOOLPULSE_SECRET","change-this-local-secret")
app.config["MAX_CONTENT_LENGTH"]=25*1024*1024
ALLOWED_RESOURCE_EXTENSIONS={"pdf","doc","docx","xls","xlsx","ppt","pptx","jpg","jpeg","png","gif","txt","csv","zip"}

def allowed_resource(filename):
    """Return True when the uploaded resource has an allowed extension."""
    name = str(filename or "").lower().strip()
    return "." in name and name.rsplit(".", 1)[1] in ALLOWED_RESOURCE_EXTENSIONS

RESOURCE_CATEGORIES=["Circular","Home Work","Time Table","Tutorials","Question Bank","Syllabus","Co-Curricular","Date Sheet","Infirmary","Communication","Examination"]

DEFAULT_EXAMS=["Unit Test 1","Unit Test 2","Half Yearly","Annual"]
DEFAULT_SUBJECTS=["English","Hindi","Mathematics","Science","Social Science","Physics","Chemistry","Biology","Computer Science","Informatics Practices","Accountancy","Economics","Business Studies"]
HEADER_ALIASES={
 "admission_no":["admission no","admission number","admission","adm no","adm no.","admission id","student admission no"],
 "name":["name","student name","student","full name","student full name"],
 "class_name":["class","class name","standard","std","grade"],
 "section":["section","sec","division","div"],
 "roll":["roll","roll no","roll no.","roll number","rollno"],
 "gender":["gender","sex"],
 "dob":["dob","date of birth","birth date","birthdate"],
 "father_name":["father","father name","fathers name","guardian","parent name"],
 "mother_name":["mother","mother name","mothers name"],
 "phone":["phone","mobile","mobile no","contact","contact no","phone number"],
 "email":["email","e-mail","mail"],
}
def db():
 c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; c.execute("PRAGMA foreign_keys=ON"); return c
def init_db():
 with db() as c:
  c.executescript("""
  CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);
  CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,username TEXT UNIQUE NOT NULL,password_hash TEXT NOT NULL,role TEXT NOT NULL,name TEXT,dob TEXT,student_id INTEGER,force_change INTEGER DEFAULT 0,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
  CREATE TABLE IF NOT EXISTS students(id INTEGER PRIMARY KEY,admission_no TEXT UNIQUE NOT NULL,name TEXT NOT NULL,class_name TEXT NOT NULL,section TEXT NOT NULL,roll INTEGER NOT NULL,gender TEXT,dob TEXT,father_name TEXT,mother_name TEXT,phone TEXT,email TEXT,created_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(class_name,section,roll));
  CREATE TABLE IF NOT EXISTS exams(id INTEGER PRIMARY KEY,name TEXT UNIQUE NOT NULL,active INTEGER DEFAULT 1,created_at TEXT DEFAULT CURRENT_TIMESTAMP);
  CREATE TABLE IF NOT EXISTS marks(student_id INTEGER,exam_id INTEGER,subject TEXT,score REAL,max_score REAL DEFAULT 100,PRIMARY KEY(student_id,exam_id,subject),FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE,FOREIGN KEY(exam_id) REFERENCES exams(id) ON DELETE CASCADE);
  CREATE TABLE IF NOT EXISTS attendance(id INTEGER PRIMARY KEY,student_id INTEGER,year INTEGER,month INTEGER,working_days INTEGER,present_days INTEGER,UNIQUE(student_id,year,month),FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE);
  CREATE TABLE IF NOT EXISTS daily_attendance(id INTEGER PRIMARY KEY,student_id INTEGER NOT NULL,attendance_date TEXT NOT NULL,status TEXT NOT NULL CHECK(status IN ('present','absent')),marked_by INTEGER,created_at TEXT DEFAULT CURRENT_TIMESTAMP,UNIQUE(student_id,attendance_date),FOREIGN KEY(student_id) REFERENCES students(id) ON DELETE CASCADE,FOREIGN KEY(marked_by) REFERENCES users(id) ON DELETE SET NULL);
  CREATE TABLE IF NOT EXISTS resources(id INTEGER PRIMARY KEY,title TEXT NOT NULL,category TEXT NOT NULL,description TEXT,class_name TEXT DEFAULT '',section TEXT DEFAULT '',original_filename TEXT NOT NULL,stored_filename TEXT UNIQUE NOT NULL,mime_type TEXT,uploaded_by INTEGER,created_at TEXT DEFAULT CURRENT_TIMESTAMP,FOREIGN KEY(uploaded_by) REFERENCES users(id) ON DELETE SET NULL);
  """)
  if not c.execute("SELECT 1 FROM exams LIMIT 1").fetchone():
   c.executemany("INSERT INTO exams(name) VALUES(?)",[(x,) for x in DEFAULT_EXAMS])
  c.execute("INSERT OR IGNORE INTO settings(key,value) VALUES('subjects',?)",(",".join(DEFAULT_SUBJECTS),))
init_db()
# Teachers keep their current password; changing it is optional.
with db() as _c:
 _c.execute("UPDATE users SET force_change=0 WHERE role='teacher'")

@app.errorhandler(413)
def request_too_large(_):
 return jsonify(error="The file is larger than the 25 MB upload limit."),413

def norm(v):
 return re.sub(r"[^a-z0-9]+"," ",str(v or "").strip().lower()).strip()
def get_subjects():
 with db() as c:
  r=c.execute("SELECT value FROM settings WHERE key='subjects'").fetchone()
 return [x.strip() for x in (r["value"] if r else ",".join(DEFAULT_SUBJECTS)).split(",") if x.strip()]
def set_subjects(items):
 items=list(dict.fromkeys([str(x).strip() for x in items if str(x).strip()]))
 with db() as c:c.execute("INSERT INTO settings(key,value) VALUES('subjects',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(",".join(items),))
def user():
 if not session.get("uid"): return None
 with db() as c:return c.execute("SELECT * FROM users WHERE id=?",(session["uid"],)).fetchone()
def auth(role=None):
 def deco(fn):
  @wraps(fn)
  def wrap(*a,**kw):
   u=user()
   if not u:return jsonify(error="Please log in."),401
   if role and u["role"]!=role:return jsonify(error="Access denied."),403
   return fn(*a,**kw)
  return wrap
 return deco
def jd():
 return request.get_json(silent=True) or {}
def code_clean(v):
 return re.sub(r"\s+","",str(v or "").strip())
def student_dict(r):
 d=dict(r)
 d["id"]=int(d["id"]); d["roll"]=int(d["roll"])
 return d

@app.get("/")
def home(): return render_template("index.html")
@app.get("/embed")
def embed(): return render_template("embed.html")
@app.get("/api/status")
def status():
 with db() as c: exists=c.execute("SELECT 1 FROM users WHERE role='principal' LIMIT 1").fetchone()
 return jsonify(setup=not bool(exists))
@app.post("/api/setup")
def setup():
 d=jd(); name=str(d.get("name","Principal")).strip(); u=code_clean(d.get("username","principal")).lower(); p=str(d.get("password","")); confirm=str(d.get("confirm_password",""))
 if len(u)<4 or " " in u or len(p)<8:return jsonify(error="Use a username without spaces/caps and a password of at least 8 characters."),400
 if p!=confirm:return jsonify(error="Passwords do not match."),400
 with db() as c:
  if c.execute("SELECT 1 FROM users WHERE role='principal' LIMIT 1").fetchone():return jsonify(error="Principal account already exists."),409
  try:c.execute("INSERT INTO users(username,password_hash,role,name,force_change) VALUES(?,?,?,?,0)",(u,generate_password_hash(p),"principal",name))
  except sqlite3.IntegrityError:return jsonify(error="Username already exists."),409
 return jsonify(ok=True)

@app.post("/api/login")
def login():
 d=jd()
 raw_username=code_clean(d.get("username",""))
 p=str(d.get("password",""))
 role=str(d.get("role","student")).strip().lower()
 if role not in {"principal","teacher","student"}:return jsonify(error="Choose Principal, Teacher or Student login."),400
 with db() as c:
  # Student admission numbers are intentionally case-insensitive at login.
  # This fixes accounts such as ADM1001 when the login form normalizes input.
  r=c.execute("SELECT * FROM users WHERE LOWER(username)=LOWER(?) AND role=?",(raw_username,role)).fetchone()
  # Older databases may contain students without a corresponding user account.
  # Create the missing account automatically using Admission No. as the initial password.
  if role=="student" and not r and raw_username:
   st=c.execute("SELECT * FROM students WHERE LOWER(admission_no)=LOWER(?)",(raw_username,)).fetchone()
   if st:
    existing=c.execute("SELECT id FROM users WHERE student_id=?",(st["id"],)).fetchone()
    if not existing:
     c.execute("INSERT INTO users(username,password_hash,role,name,student_id,force_change) VALUES(?,?, 'student',?,?,0)",(st["admission_no"],generate_password_hash(st["admission_no"]),st["name"],st["id"]))
     r=c.execute("SELECT * FROM users WHERE student_id=? AND role='student'",(st["id"],)).fetchone()
 if not r or not check_password_hash(r["password_hash"],p):
  label="Admission No." if role=="student" else "User ID"
  return jsonify(error=f"Invalid {label} or password for {role.title()} login."),401
 session.clear();session["uid"]=r["id"]
 return jsonify(ok=True,role=r["role"],force_change=bool(r["force_change"]))
@app.post("/api/logout")
def logout():session.clear();return jsonify(ok=True)
@app.get("/api/me")
def me():
 u=user()
 if not u:return jsonify(user=None)
 return jsonify(user={"id":u["id"],"username":u["username"],"role":u["role"],"name":u["name"],"student_id":u["student_id"],"force_change":bool(u["force_change"])})

@app.post("/api/change-password")
@auth()
def change_password():
 d=jd(); old=str(d.get("old","")); new=str(d.get("new",""))
 u=user()
 if not check_password_hash(u["password_hash"],old):return jsonify(error="Current password is incorrect."),400
 if len(new)<8:return jsonify(error="New password must be at least 8 characters."),400
 with db() as c:c.execute("UPDATE users SET password_hash=?,force_change=0 WHERE id=?",(generate_password_hash(new),u["id"]))
 return jsonify(ok=True)

@app.get("/api/config")
@auth()
def config():
 with db() as c:ex=c.execute("SELECT id,name FROM exams WHERE active=1 ORDER BY id").fetchall()
 return jsonify(exams=[dict(x) for x in ex],subjects=get_subjects(),classes=[str(i) for i in range(1,13)],resource_categories=RESOURCE_CATEGORIES)

@app.post("/api/exams")
@auth()
def add_exam():
 n=str(jd().get("name","")).strip()
 if len(n)<2:return jsonify(error="Enter an exam name."),400
 try:
  with db() as c:c.execute("INSERT INTO exams(name) VALUES(?)",(n,))
 except sqlite3.IntegrityError:return jsonify(error="That exam already exists."),409
 return jsonify(ok=True)
@app.delete("/api/exams/<int:eid>")
@auth("principal")
def del_exam(eid):
 with db() as c:
  c.execute("DELETE FROM marks WHERE exam_id=?",(eid,));c.execute("DELETE FROM exams WHERE id=?",(eid,))
 return jsonify(ok=True)

@app.post("/api/teachers")
@auth("principal")
def create_teacher():
 d=jd();name=str(d.get("name","")).strip();dob=str(d.get("dob","")).strip();username=code_clean(d.get("username","")).lower()
 if not name or not dob or not username or " " in username:return jsonify(error="Name, DOB and a lowercase username without spaces are required."),400
 initial=str(d.get("password","")).strip()
 if not initial:
  digits=re.sub(r"[^0-9]","",dob)
  first=re.sub(r"[^a-z]","",name.lower().split()[0])
  initial=(first+digits)[:40]
 if len(initial)<8:return jsonify(error="Initial password must be at least 8 characters."),400
 try:
  with db() as c:c.execute("INSERT INTO users(username,password_hash,role,name,dob,force_change) VALUES(?,?, 'teacher',?,?,0)",(username,generate_password_hash(initial),name,dob))
 except sqlite3.IntegrityError:return jsonify(error="Teacher username already exists."),409
 return jsonify(ok=True,username=username,initial_password=initial)

@app.get("/api/teachers")
@auth("principal")
def teachers():
 with db() as c:r=c.execute("SELECT id,username,name,dob,created_at FROM users WHERE role='teacher' ORDER BY name").fetchall()
 return jsonify([dict(x) for x in r])

def create_student_record(d, c):
 admission=code_clean(d.get("admission_no",""))
 name=str(d.get("name","")).strip()
 cl=str(d.get("class_name","")).strip()
 sec=str(d.get("section","")).strip().upper()
 try:roll=int(d.get("roll"))
 except:return None,"Roll number must be numeric."
 if not admission or not name or cl not in [str(i) for i in range(1,13)] or not sec or len(sec)>4 or roll<1:return None,"Admission no, name, class, section and roll are required."
 try:
  cur=c.execute("""INSERT INTO students(admission_no,name,class_name,section,roll,gender,dob,father_name,mother_name,phone,email)
  VALUES(?,?,?,?,?,?,?,?,?,?,?)""",(admission,name,cl,sec,roll,d.get("gender",""),d.get("dob",""),d.get("father_name",""),d.get("mother_name",""),d.get("phone",""),d.get("email","")))
 except sqlite3.IntegrityError as e:
  return None,"Duplicate admission number or duplicate class-section-roll."
 sid=cur.lastrowid
 # Admission number is BOTH username and initial password, per requested workflow.
 c.execute("INSERT INTO users(username,password_hash,role,name,student_id,force_change) VALUES(?,?, 'student',?,?,0)",(admission,generate_password_hash(admission),name,sid))
 return sid,None

@app.get("/api/students")
@auth()
def students():
 u=user()
 with db() as c:
  if u["role"]=="student":r=c.execute("SELECT * FROM students WHERE id=?",(u["student_id"],)).fetchall()
  else:r=c.execute("SELECT * FROM students ORDER BY CAST(class_name AS INTEGER),section,roll").fetchall()
 return jsonify([student_dict(x) for x in r])

@app.post("/api/students")
@auth("teacher")
def add_student():
 d=jd()
 try:
  with db() as c:
   sid,err=create_student_record(d,c)
   if err:return jsonify(error=err),400
 except Exception as e:return jsonify(error=str(e)),400
 return jsonify(ok=True,student_id=sid,admission_no=code_clean(d.get("admission_no"))),201

@app.put("/api/students/<int:sid>")
@auth("teacher")
def edit_student(sid):
 d=jd(); admission=code_clean(d.get("admission_no","")); name=str(d.get("name"," ")).strip(); cl=str(d.get("class_name"," ")).strip(); sec=str(d.get("section"," ")).strip().upper()
 try: roll=int(d.get("roll"))
 except: return jsonify(error="Roll number must be numeric."),400
 if not admission or not name or cl not in [str(i) for i in range(1,13)] or not sec or len(sec)>4 or roll<1:return jsonify(error="Admission no, name, class, section and roll are required."),400
 with db() as c:
  if not c.execute("SELECT id FROM students WHERE id=?",(sid,)).fetchone(): return jsonify(error="Student not found."),404
  try:
   c.execute("""UPDATE students SET admission_no=?,name=?,class_name=?,section=?,roll=?,gender=?,dob=?,father_name=?,mother_name=?,phone=?,email=? WHERE id=?""",(admission,name,cl,sec,roll,d.get("gender",""),d.get("dob",""),d.get("father_name",""),d.get("mother_name",""),d.get("phone",""),d.get("email",""),sid))
   c.execute("UPDATE users SET username=?,name=? WHERE student_id=?",(admission,name,sid))
  except sqlite3.IntegrityError:return jsonify(error="Admission number or class-section-roll already exists."),409
 return jsonify(ok=True)

@app.delete("/api/students/<int:sid>")
@auth("teacher")
def delete_student(sid):
 with db() as c:
  c.execute("DELETE FROM users WHERE student_id=?",(sid,));cur=c.execute("DELETE FROM students WHERE id=?",(sid,))
 return jsonify(ok=bool(cur.rowcount))

@app.get("/api/marks/<int:sid>")
@auth()
def marks(sid):
 u=user()
 if u["role"]=="student" and u["student_id"]!=sid:return jsonify(error="Access denied."),403
 with db() as c:r=c.execute("""SELECT m.*,e.name exam FROM marks m JOIN exams e ON e.id=m.exam_id WHERE m.student_id=? ORDER BY e.id,m.subject""",(sid,)).fetchall()
 return jsonify([dict(x) for x in r])
@app.post("/api/marks")
@auth("teacher")
def save_mark():
 d=jd()
 try:sid=int(d["student_id"]);eid=int(d["exam_id"]);score=float(d["score"]);mx=float(d.get("max_score",100));sub=str(d["subject"]).strip()
 except:return jsonify(error="Invalid marks data."),400
 if sub not in get_subjects() or score<0 or mx<=0 or score>mx:return jsonify(error="Invalid subject or marks."),400
 with db() as c:c.execute("""INSERT INTO marks VALUES(?,?,?,?,?) ON CONFLICT(student_id,exam_id,subject) DO UPDATE SET score=excluded.score,max_score=excluded.max_score""",(sid,eid,sub,score,mx))
 return jsonify(ok=True)

@app.get("/api/attendance/<int:sid>")
@auth()
def get_attendance(sid):
 u=user()
 if u["role"]=="student" and u["student_id"]!=sid:return jsonify(error="Access denied."),403
 with db() as c:r=c.execute("SELECT * FROM attendance WHERE student_id=? ORDER BY year,month",(sid,)).fetchall()
 return jsonify([{**dict(x),"percentage":round((x["present_days"]/x["working_days"]*100),1) if x["working_days"] else 0} for x in r])
@app.get("/api/attendance/class")
@auth("teacher")
def class_attendance():
 cl=str(request.args.get("class_name","")); sec=str(request.args.get("section","" )).upper(); year=int(request.args.get("year",datetime.now().year)); month=int(request.args.get("month",datetime.now().month))
 with db() as c:
  rows=c.execute("""SELECT s.id,s.name,s.admission_no,s.roll,s.class_name,s.section,COALESCE(a.working_days,0) working_days,COALESCE(a.present_days,0) present_days FROM students s LEFT JOIN attendance a ON a.student_id=s.id AND a.year=? AND a.month=? WHERE s.class_name=? AND s.section=? ORDER BY s.roll""",(year,month,cl,sec)).fetchall()
 return jsonify(year=year,month=month,class_name=cl,section=sec,students=[{**dict(r),"percentage":round(r["present_days"]/r["working_days"]*100,1) if r["working_days"] else None} for r in rows])

@app.post("/api/attendance/bulk")
@auth("teacher")
def save_attendance_bulk():
 d=jd();cl=str(d.get("class_name",""));sec=str(d.get("section","" )).upper()
 try: year=int(d.get("year"));month=int(d.get("month"));wd=int(d.get("working_days"));records=d.get("records",[])
 except:return jsonify(error="Invalid attendance data."),400
 if cl not in [str(i) for i in range(1,13)] or not sec or not(1<=month<=12) or wd<1:return jsonify(error="Enter a valid class, section, month and working days."),400
 with db() as c:
  ids={int(r["id"]):int(r.get("present_days",0)) for r in records if str(r.get("id","")).isdigit()}
  valid={r["id"] for r in c.execute("SELECT id FROM students WHERE class_name=? AND section=?",(cl,sec)).fetchall()}
  if not set(ids).issubset(valid):return jsonify(error="Attendance list contains an invalid student."),400
  for sid,pd in ids.items():
   if pd<0 or pd>wd:return jsonify(error=f"Present days cannot exceed working days for student ID {sid}."),400
   c.execute("""INSERT INTO attendance(student_id,year,month,working_days,present_days) VALUES(?,?,?,?,?) ON CONFLICT(student_id,year,month) DO UPDATE SET working_days=excluded.working_days,present_days=excluded.present_days""",(sid,year,month,wd,pd))
 return jsonify(ok=True,saved=len(ids),working_days=wd)

@app.get("/api/attendance/daily/class")
@auth("teacher")
def daily_class_attendance():
 date=str(request.args.get("date",datetime.now().date().isoformat()))
 cl=str(request.args.get("class_name",""));sec=str(request.args.get("section","" )).upper()
 with db() as c:
  rows=c.execute("""SELECT s.id,s.name,s.admission_no,s.roll,s.class_name,s.section,COALESCE(d.status,'absent') status FROM students s LEFT JOIN daily_attendance d ON d.student_id=s.id AND d.attendance_date=? WHERE s.class_name=? AND s.section=? ORDER BY s.roll""",(date,cl,sec)).fetchall()
 return jsonify(date=date,class_name=cl,section=sec,students=[dict(r) for r in rows])

@app.post("/api/attendance/daily/bulk")
@auth("teacher")
def save_daily_attendance():
 d=jd();date=str(d.get("date",""));records=d.get("records",[])
 try: datetime.strptime(date,"%Y-%m-%d")
 except: return jsonify(error="Enter a valid date."),400
 if not records:return jsonify(error="No attendance records supplied."),400
 with db() as c:
  valid={r["id"] for r in c.execute("SELECT id FROM students WHERE class_name=? AND section=?",(str(d.get("class_name","")).strip(),str(d.get("section","")).upper().strip())).fetchall()}
  for r in records:
   try:sid=int(r["id"])
   except:return jsonify(error="Invalid student record."),400
   if sid not in valid or r.get("status") not in {"present","absent"}:return jsonify(error="Invalid daily attendance record."),400
   c.execute("""INSERT INTO daily_attendance(student_id,attendance_date,status,marked_by) VALUES(?,?,?,?) ON CONFLICT(student_id,attendance_date) DO UPDATE SET status=excluded.status,marked_by=excluded.marked_by""",(sid,date,r["status"],user()["id"]))
 return jsonify(ok=True,saved=len(records))

@app.get("/api/attendance/student/<int:sid>/summary")
@auth()
def student_attendance_summary(sid):
 u=user()
 if u["role"]=="student" and u["student_id"]!=sid:return jsonify(error="Access denied."),403
 with db() as c:
  daily=c.execute("SELECT attendance_date,status FROM daily_attendance WHERE student_id=? ORDER BY attendance_date DESC",(sid,)).fetchall()
  monthly=c.execute("SELECT * FROM attendance WHERE student_id=? ORDER BY year DESC,month DESC",(sid,)).fetchall()
 total_days=len(daily);present=sum(x["status"]=="present" for x in daily)
 return jsonify(total_days=total_days,present_days=present,absent_days=total_days-present,percentage=round(present/total_days*100,1) if total_days else None,absences=[dict(x) for x in daily if x["status"]=="absent"],daily=[dict(x) for x in daily],monthly=[{**dict(x),"percentage":round(x["present_days"]/x["working_days"]*100,1) if x["working_days"] else 0} for x in monthly])

@app.post("/api/attendance")
@auth("teacher")
def save_attendance():
 d=jd()
 try:sid=int(d["student_id"]);year=int(d["year"]);month=int(d["month"]);wd=int(d["working_days"]);pd=int(d["present_days"])
 except:return jsonify(error="Invalid attendance data."),400
 if not(1<=month<=12 and 0<=pd<=wd and wd>0):return jsonify(error="Present days cannot exceed working days."),400
 with db() as c:c.execute("""INSERT INTO attendance(student_id,year,month,working_days,present_days) VALUES(?,?,?,?,?)
 ON CONFLICT(student_id,year,month) DO UPDATE SET working_days=excluded.working_days,present_days=excluded.present_days""",(sid,year,month,wd,pd))
 return jsonify(ok=True,percentage=round(pd/wd*100,1))

def detect_headers(headers):
 normalized={norm(h):h for h in headers}
 mapping={}
 for field,aliases in HEADER_ALIASES.items():
  targets=[norm(x) for x in aliases]
  hit=next((normalized[t] for t in targets if t in normalized),None)
  if not hit:
   for nh,orig in normalized.items():
    if any(t in nh or nh in t for t in targets if t):hit=orig;break
  mapping[field]=hit
 return mapping

@app.post("/api/import/preview")
@auth("teacher")
def import_preview():
 f=request.files.get("file")
 if not f or not f.filename.lower().endswith(".xlsx"):return jsonify(error="Upload an .xlsx Excel file."),400
 path=UPLOADS/(secrets.token_hex(8)+".xlsx");f.save(path)
 wb=None
 try:
  wb=load_workbook(path,read_only=True,data_only=True)
  ws=wb.active;rows=ws.iter_rows(values_only=True)
  headers=next(rows);mapping=detect_headers(headers)
  sample=[]
  for vals in rows:
   if not any(v is not None and str(v).strip() for v in vals):continue
   raw={headers[i]:vals[i] for i in range(len(headers))}
   sample.append({k:(raw.get(v) if v else "") for k,v in mapping.items()})
   if len(sample)>=8:break
  return jsonify(token=path.name,headers=list(headers),mapping=mapping,sample=sample)
 except Exception as e:
  return jsonify(error=f"Could not read Excel file: {e}"),400
 finally:
  # Windows keeps a read-only workbook file handle open until close().
  # Always close it before the uploaded file is reused/deleted.
  if wb is not None:
   try: wb.close()
   except Exception: pass

@app.post("/api/import/commit")
@auth("teacher")
def import_commit():
 token=code_clean(jd().get("token",""))
 path=UPLOADS/token
 if not token or not path.exists():return jsonify(error="Upload preview expired. Upload the file again."),400
 wb=None
 try:
  wb=load_workbook(path,read_only=True,data_only=True);ws=wb.active;rows=ws.iter_rows(values_only=True);headers=list(next(rows));mapping=jd().get("mapping") or detect_headers(headers)
  created=skipped=0;errors=[]
  with db() as c:
   for line,vals in enumerate(rows,start=2):
    if not any(v is not None and str(v).strip() for v in vals):continue
    raw={headers[i]:vals[i] for i in range(len(headers))}
    d={k:raw.get(v,"") for k,v in mapping.items() if v}
    if not d.get("admission_no") or not d.get("name"):skipped+=1;errors.append(f"Row {line}: admission no/name missing");continue
    sid,err=create_student_record(d,c)
    if err:skipped+=1;errors.append(f"Row {line}: {err}")
    else:created+=1
  return jsonify(created=created,skipped=skipped,errors=errors[:20])
 except Exception as e:return jsonify(error=str(e)),400
 finally:
  if wb is not None:
   try: wb.close()
   except Exception: pass
  # Delete only after openpyxl has released the Windows file handle.
  try: path.unlink(missing_ok=True)
  except PermissionError: pass


@app.get("/api/resources")
@auth()
def resources():
 u=user(); category=str(request.args.get("category","")).strip()
 with db() as c:
  if u["role"]=="student":
   s=c.execute("SELECT class_name,section FROM students WHERE id=?",(u["student_id"],)).fetchone()
   if not s:return jsonify([])
   q="SELECT r.*,u.name uploader FROM resources r LEFT JOIN users u ON u.id=r.uploaded_by WHERE (r.class_name='' OR (r.class_name=? AND (r.section='' OR r.section=?)))"
   args=[s["class_name"],s["section"]]
   if category:q+=" AND r.category=?";args.append(category)
   q+=" ORDER BY datetime(r.created_at) DESC"
   rows=c.execute(q,args).fetchall()
  else:
   q="SELECT r.*,u.name uploader FROM resources r LEFT JOIN users u ON u.id=r.uploaded_by";args=[]
   if category:q+=" WHERE r.category=?";args.append(category)
   q+=" ORDER BY datetime(r.created_at) DESC"
   rows=c.execute(q,args).fetchall()
 return jsonify([{**dict(r),"download_url":f"/api/resources/{r['id']}/download"} for r in rows])

@app.post("/api/resources")
@auth()
def upload_resource():
 u=user()
 if u["role"] not in {"principal","teacher"}:
  return jsonify(ok=False,error="Only teachers and principal can upload academic resources."),403

 title=str(request.form.get("title","")).strip()
 category=str(request.form.get("category","")).strip()
 description=str(request.form.get("description","")).strip()
 class_name=str(request.form.get("class_name","")).strip()
 section=str(request.form.get("section","")).strip().upper()
 f=request.files.get("file")

 if category not in RESOURCE_CATEGORIES:
  return jsonify(ok=False,error="Choose a valid resource category."),400
 if not title:
  return jsonify(ok=False,error="Enter a resource title."),400
 if not f or not f.filename:
  return jsonify(ok=False,error="Choose a file to upload."),400
 if not allowed_resource(f.filename):
  return jsonify(ok=False,error="File type is not allowed. Use PDF, Word, Excel, PowerPoint, images, TXT, CSV or ZIP."),400
 if class_name and class_name not in [str(i) for i in range(1,13)]:
  return jsonify(ok=False,error="Invalid class."),400
 if section and len(section)>4:
  return jsonify(ok=False,error="Invalid section."),400

 original_name=Path(f.filename).name
 ext=original_name.rsplit(".",1)[1].lower()
 import uuid
 stored=f"{uuid.uuid4().hex}.{ext}"
 UPLOADS.mkdir(parents=True, exist_ok=True)
 path=UPLOADS/stored

 try:
  f.save(str(path))
  if not path.exists() or path.stat().st_size == 0:
   raise OSError("The uploaded file could not be written to the uploads folder.")

  with db() as c:
   cur=c.execute(
    "INSERT INTO resources(title,category,description,class_name,section,original_filename,stored_filename,mime_type,uploaded_by) VALUES(?,?,?,?,?,?,?,?,?)",
    (title,category,description,class_name,section,original_name,stored,f.mimetype or "application/octet-stream",u["id"])
   )

  return jsonify(
   ok=True,
   id=cur.lastrowid,
   filename=original_name,
   title=title,
   category=category,
   class_name=class_name,
   section=section
  ),201

 except Exception as exc:
  try:
   path.unlink(missing_ok=True)
  except Exception:
   pass
  app.logger.exception("Resource upload failed")
  return jsonify(ok=False,error=f"Upload failed: {exc}"),500

@app.delete("/api/resources/<int:rid>")
@auth()
def delete_resource(rid):
 u=user()
 with db() as c:r=c.execute("SELECT * FROM resources WHERE id=?",(rid,)).fetchone()
 if not r:return jsonify(error="Resource not found."),404
 if u["role"]!="principal" and r["uploaded_by"]!=u["id"]:return jsonify(error="You can only remove resources you uploaded."),403
 try:(UPLOADS/r["stored_filename"]).unlink(missing_ok=True)
 except Exception:pass
 with db() as c:c.execute("DELETE FROM resources WHERE id=?",(rid,))
 return jsonify(ok=True)

@app.get("/api/resources/<int:rid>/download")
@auth()
def download_resource(rid):
 u=user()
 with db() as c:
  if u["role"]=="student":
   r=c.execute("SELECT r.*,s.class_name student_class,s.section student_section FROM resources r JOIN students s ON s.id=? WHERE r.id=?",(u["student_id"],rid)).fetchone()
  else:r=c.execute("SELECT * FROM resources WHERE id=?",(rid,)).fetchone()
 if not r:return jsonify(error="Resource not found."),404
 if u["role"]=="student":
  if r["class_name"] and r["class_name"]!=r["student_class"]:return jsonify(error="This resource is not available to your class."),403
  if r["section"] and r["section"]!=r["student_section"]:return jsonify(error="This resource is not available to your section."),403
 return send_from_directory(UPLOADS,r["stored_filename"],as_attachment=True,download_name=r["original_filename"])

@app.get("/api/analytics")
@auth()
def analytics():
 u=user()
 with db() as c:
  if u["role"]=="student":students=c.execute("SELECT * FROM students WHERE id=?",(u["student_id"],)).fetchall()
  else:students=c.execute("SELECT * FROM students").fetchall()
  ids=[x["id"] for x in students]
  if not ids:return jsonify(students=0,classes=[],average=None,attendance=None,subject={},exam={})
  marks=c.execute("SELECT m.*,e.name exam FROM marks m JOIN exams e ON e.id=m.exam_id WHERE m.student_id IN (%s)"%(",".join("?"*len(ids))),ids).fetchall()
 vals=[m["score"]/m["max_score"]*100 for m in marks if m["max_score"]]
 subj={}
 for m in marks:
  subj.setdefault(m["subject"],[]).append(m["score"]/m["max_score"]*100)
 exam={}
 for m in marks:
  exam.setdefault(m["exam"],[]).append(m["score"]/m["max_score"]*100)
 at=c.execute("SELECT working_days,present_days FROM attendance WHERE student_id IN (%s)"%(",".join("?"*len(ids))),ids).fetchall()
 att=sum(x["present_days"] for x in at)/sum(x["working_days"] for x in at)*100 if sum(x["working_days"] for x in at) else None
 return jsonify(students=len(students),classes=sorted({f'{x["class_name"]}-{x["section"]}' for x in students}),average=round(sum(vals)/len(vals),1) if vals else None,attendance=round(att,1) if att is not None else None,subject={k:round(sum(v)/len(v),1) for k,v in subj.items()},exam={k:round(sum(v)/len(v),1) for k,v in exam.items()})


def _report_student_or_forbidden(sid):
    u=user()
    if u["role"]=="student" and int(u["student_id"] or 0)!=int(sid):
        return None, ("Access denied.",403)
    with db() as c:
        st=c.execute("SELECT * FROM students WHERE id=?",(sid,)).fetchone()
    if not st:
        return None, ("Student not found.",404)
    return st, None

def _grade(pct):
    try: p=float(pct)
    except: return "—"
    if p >= 91: return "A1"
    if p >= 81: return "A2"
    if p >= 71: return "B1"
    if p >= 61: return "B2"
    if p >= 51: return "C1"
    if p >= 41: return "C2"
    if p >= 33: return "D"
    return "E"

@app.get("/reports/<kind>/<int:sid>")
@auth()
def print_report(kind,sid):
    if kind not in {"student","attendance","academic","marksheet"}:
        return "Unknown report.",404
    st,err=_report_student_or_forbidden(sid)
    if err:return err
    with db() as c:
        exams=c.execute("SELECT * FROM exams WHERE active=1 ORDER BY id").fetchall()
        marks=c.execute("""SELECT m.*,e.name exam_name FROM marks m JOIN exams e ON e.id=m.exam_id
                           WHERE m.student_id=? ORDER BY e.id,m.subject""",(sid,)).fetchall()
        daily=c.execute("""SELECT attendance_date,status FROM daily_attendance
                           WHERE student_id=? ORDER BY attendance_date""",(sid,)).fetchall()
        monthly=c.execute("""SELECT year,month,working_days,present_days FROM attendance
                             WHERE student_id=? ORDER BY year,month""",(sid,)).fetchall()
    exam_map={}
    for r in marks:
        exam_map.setdefault(r["exam_name"],[]).append(dict(r))
    total_max=sum(float(r["max_score"] or 0) for r in marks)
    total_score=sum(float(r["score"] or 0) for r in marks)
    overall_pct=(total_score/total_max*100) if total_max else 0
    present=sum(1 for r in daily if r["status"]=="present")
    absent=sum(1 for r in daily if r["status"]=="absent")
    if daily:
        att_pct=present/len(daily)*100
    else:
        w=sum(int(r["working_days"] or 0) for r in monthly)
        p=sum(int(r["present_days"] or 0) for r in monthly)
        att_pct=(p/w*100) if w else 0
    subjects=get_subjects()
    return render_template("print_report.html",kind=kind,student=dict(st),exams=[dict(x) for x in exams],
                           exam_map=exam_map,subjects=subjects,marks=[dict(x) for x in marks],
                           daily=[dict(x) for x in daily],monthly=[dict(x) for x in monthly],
                           total_max=total_max,total_score=total_score,overall_pct=overall_pct,
                           grade=_grade(overall_pct),present=present,absent=absent,att_pct=att_pct,
                           generated=datetime.now().strftime("%d %B %Y, %I:%M %p"))

@app.get("/reports/class/<class_name>/<section>")
@auth()
def print_class_report(class_name,section):
    if user()["role"]=="student": return "Access denied.",403
    with db() as c:
        sts=c.execute("""SELECT * FROM students WHERE class_name=? AND section=?
                         ORDER BY roll""",(class_name,section.upper())).fetchall()
        rows=[]
        for st in sts:
            sid=st["id"]
            marks=c.execute("""SELECT m.score,m.max_score FROM marks m WHERE m.student_id=?""",(sid,)).fetchall()
            mx=sum(float(x["max_score"] or 0) for x in marks); sc=sum(float(x["score"] or 0) for x in marks)
            pct=sc/mx*100 if mx else 0
            daily=c.execute("""SELECT status FROM daily_attendance WHERE student_id=?""",(sid,)).fetchall()
            if daily:
                ap=sum(x["status"]=="present" for x in daily)/len(daily)*100
            else:
                mon=c.execute("""SELECT working_days,present_days FROM attendance WHERE student_id=?""",(sid,)).fetchall()
                w=sum(int(x["working_days"] or 0) for x in mon);p=sum(int(x["present_days"] or 0) for x in mon)
                ap=p/w*100 if w else 0
            rows.append({"roll":st["roll"],"name":st["name"],"admission_no":st["admission_no"],"score":sc,"max":mx,"pct":pct,"grade":_grade(pct),"att":ap})
    return render_template("print_class_report.html",class_name=class_name,section=section.upper(),rows=rows,
                           generated=datetime.now().strftime("%d %B %Y, %I:%M %p"))

@app.get("/api/site-info")
def site_info():
 return jsonify(name="EduTrack Academic",version="2.0",embed_url="/embed")
if __name__=="__main__":app.run(host="127.0.0.1",port=5000,debug=True)
